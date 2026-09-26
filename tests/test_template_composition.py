from __future__ import annotations

import json
import sqlite3
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app.config import settings
from app.core.permissions import require_user
from app.core.ppt import split_pptx_to_single_pages
from app.db import init_db, now_iso
from app.routers import resource_import, tasks, templates
from app.routers.dependencies import db_dep, db_read_dep
from app.services.files import _init_allowed_file_dirs
from app.services.resource_import.remote_fonts import sha256_file
from app.services.resource_import.rendering import RESOURCE_IMPORT_RENDERER_VERSION
from app.services.resource_import.sessions import _load_resource_import_session_file, _write_resource_import_session


class TemplateCompositionTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="slideflow-template-compose-")
        root = Path(self.temporary.name)
        paths = {
            "root_dir": root,
            "data_dir": root / "data",
            "db_dir": root / "data" / "db",
            "assets_dir": root / "data" / "assets",
            "resources_dir": root / "data" / "assets" / "resources",
            "templates_dir": root / "data" / "assets" / "templates",
            "fonts_dir": root / "data" / "assets" / "fonts",
            "thumbs_dir": root / "data" / "assets" / "thumbs",
            "downloads_dir": root / "data" / "assets" / "downloads",
            "log_dir": root / "data" / "logs",
            "db_path": root / "data" / "db" / "slide_flow.db",
        }
        self.patchers = [patch.object(settings, name, value) for name, value in paths.items()]
        self.patchers.append(patch.object(settings, "storage_backend", "local"))
        for patcher in self.patchers:
            patcher.start()

        init_db()
        _init_allowed_file_dirs()
        self.db = sqlite3.connect(settings.db_path, check_same_thread=False)
        self.db.row_factory = sqlite3.Row
        self.db.execute("PRAGMA foreign_keys = ON")
        self.alice = self.insert_user("alice")
        self.bob = self.insert_user("bob")
        self.current_user = self.alice

        app = FastAPI()
        app.include_router(templates.router)
        app.include_router(resource_import.router)
        app.include_router(tasks.router)

        def override_user():
            return self.current_user

        def override_db():
            yield self.db

        app.dependency_overrides[require_user] = override_user
        app.dependency_overrides[db_dep] = override_db
        app.dependency_overrides[db_read_dep] = override_db
        self.client = TestClient(app)

    def tearDown(self):
        self.client.close()
        self.db.close()
        for patcher in reversed(self.patchers):
            patcher.stop()
        _init_allowed_file_dirs()
        self.temporary.cleanup()

    def insert_user(self, username: str) -> sqlite3.Row:
        timestamp = now_iso()
        user_id = int(
            self.db.execute(
                """
                INSERT INTO users (
                    name, username, username_key, password_hash, role, created_at, updated_at
                ) VALUES (?, ?, ?, 'not-used', 'user', ?, ?)
                """,
                (username.title(), username, username, timestamp, timestamp),
            ).lastrowid
        )
        self.db.commit()
        row = self.db.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
        assert row is not None
        return row

    def insert_template(
        self,
        *,
        name: str,
        owner: sqlite3.Row,
        text: str,
        ratio: str = "16:9",
        visibility_scope: str = "public",
    ) -> int:
        settings.templates_dir.mkdir(parents=True, exist_ok=True)
        path = settings.templates_dir / (name + ".pptx")
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        box.text_frame.text = text
        presentation.save(path)

        timestamp = now_iso()
        template_id = int(
            self.db.execute(
                """
                INSERT INTO templates (
                    name, series, subject, platform, ratio, template_type,
                    office_file_name, office_path, png_path, font_names, missing_fonts,
                    subject_order, series_order, sort_order,
                    visibility_scope, management_scope, owner_id, created_at, updated_at
                ) VALUES (?, '系列', '主体', 'wps', ?, 'content', ?, ?, NULL, '[]', '[]',
                          10, 10, 10, ?, 'private', ?, ?, ?)
                """,
                (
                    name,
                    ratio,
                    path.name,
                    str(path),
                    visibility_scope,
                    int(owner["id"]),
                    timestamp,
                    timestamp,
                ),
            ).lastrowid
        )
        self.db.commit()
        return template_id

    def test_combines_selected_templates_in_requested_order(self):
        first_id = self.insert_template(name="第一页", owner=self.alice, text="FIRST")
        second_id = self.insert_template(name="第二页", owner=self.alice, text="SECOND")

        response = self.client.post(
            "/api/templates/compose-download",
            json={"template_ids": [second_id, first_id]},
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(
            response.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )

        output = settings.downloads_dir / "combined.pptx"
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_bytes(response.content)
        presentation = Presentation(output)
        slide_texts = [
            " ".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))
            for slide in presentation.slides
        ]
        self.assertEqual(len(slide_texts), 2)
        self.assertIn("SECOND", slide_texts[0])
        self.assertIn("FIRST", slide_texts[1])

    def test_rejects_mixed_ratios_and_invisible_templates(self):
        wide_id = self.insert_template(name="宽屏", owner=self.alice, text="WIDE")
        standard_id = self.insert_template(
            name="四比三",
            owner=self.alice,
            text="STANDARD",
            ratio="4:3",
        )
        mixed = self.client.post(
            "/api/templates/compose-download",
            json={"template_ids": [wide_id, standard_id]},
        )
        self.assertEqual(mixed.status_code, 400, mixed.text)
        self.assertIn("相同比例", mixed.json()["detail"])

        private_id = self.insert_template(
            name="他人私有",
            owner=self.bob,
            text="PRIVATE",
            visibility_scope="private",
        )
        forbidden = self.client.post(
            "/api/templates/compose-download",
            json={"template_ids": [wide_id, private_id]},
        )
        self.assertEqual(forbidden.status_code, 403, forbidden.text)

    def test_commits_reviewed_ppt_as_ordered_template_series(self):
        self.db.execute("UPDATE users SET role = 'admin' WHERE id = ?", (int(self.alice["id"]),))
        self.db.commit()
        self.current_user = self.db.execute("SELECT * FROM users WHERE id = ?", (int(self.alice["id"]),)).fetchone()

        session_id = "a" * 32
        session_dir = (settings.assets_dir / ".resource_imports" / session_id).resolve()
        session_dir.mkdir(parents=True)
        source = session_dir / "source.pptx"
        presentation = Presentation()
        for text in ("FIRST", "SECOND"):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
            box.text_frame.text = text
        presentation.save(source)
        split_files = split_pptx_to_single_pages(source, session_dir / "split")
        preview_paths = []
        for index, color in enumerate(((255, 0, 0), (0, 0, 255)), start=1):
            preview = session_dir / f"preview_{index}.png"
            Image.new("RGB", (320, 180), color).save(preview)
            preview_paths.append(preview)
        session = {
            "session_id": session_id,
            "owner_id": int(self.alice["id"]),
            "mode": "ppt",
            "temp_dir": str(session_dir),
            "source_path": str(source),
            "preview_paths": [str(path) for path in preview_paths],
            "preview_hashes": [sha256_file(path) for path in preview_paths],
            "split_paths": [str(path) for path in split_files],
            "split_hashes": [sha256_file(path) for path in split_files],
            "rendered_source_sha256": sha256_file(source),
            "slide_count": 2,
            "fonts": [],
            "missing_fonts": [],
            "preview_status": "ready",
            "renderer_version": RESOURCE_IMPORT_RENDERER_VERSION,
            "expires_at": time.time() + 3600,
        }
        _write_resource_import_session(session)

        response = self.client.post(
            f"/api/resource-import/{session_id}/commit",
            json={
                "series": "秋季发布会",
                "subject": "品牌",
                "platform": "wps",
                "ratio": "16:9",
                "template_type": "content",
                "visibility_scope": "public",
                "management_scope": "private",
                "visible_user_ids": [],
                "manage_user_ids": [],
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["created"], 2)
        rows = self.db.execute(
            "SELECT name, series, subject, sort_order, office_path, png_path FROM templates ORDER BY sort_order"
        ).fetchall()
        self.assertEqual([row["series"] for row in rows], ["秋季发布会", "秋季发布会"])
        self.assertEqual([row["subject"] for row in rows], ["品牌", "品牌"])
        self.assertEqual([row["sort_order"] for row in rows], [10, 20])
        self.assertTrue(rows[0]["name"].endswith("-01"))
        self.assertTrue(rows[1]["name"].endswith("-02"))
        for row in rows:
            self.assertTrue(settings.abs_path(row["office_path"]).is_file())
            self.assertTrue(settings.abs_path(row["png_path"]).is_file())

        replay = self.client.post(
            f"/api/resource-import/{session_id}/commit",
            json={"series": "ignored"},
        )
        self.assertEqual(replay.status_code, 200, replay.text)
        self.assertEqual(replay.json()["template_ids"], response.json()["template_ids"])
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM templates").fetchone()[0], 2)

    def test_template_upload_uses_shared_durable_import_task(self):
        self.db.execute("UPDATE users SET role = 'admin' WHERE id = ?", (int(self.alice["id"]),))
        self.db.commit()
        self.current_user = self.db.execute("SELECT * FROM users WHERE id = ?", (int(self.alice["id"]),)).fetchone()
        source = settings.data_dir / "template-series.pptx"
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])
        presentation.save(source)

        with patch.object(tasks, "schedule_resource_import_task") as schedule:
            response = self.client.post(
                "/api/tasks/split-import",
                data={
                    "name_prefix": "品牌系列",
                    "subject": "品牌",
                    "tags": "",
                    "secrecy_level": "public",
                    "status": "active",
                    "visibility_scope": "public",
                    "visible_user_ids": "[]",
                    "visible_user_tags": "[]",
                    "management_scope": "private",
                    "manage_user_ids": "[]",
                    "manage_user_tags": "[]",
                    "remark_html": "",
                    "import_target": "templates",
                    "series": "品牌系列",
                    "platform": "wps",
                    "ratio": "16:9",
                    "template_type": "content",
                },
                files=[
                    ("ppt_file", (source.name, source.read_bytes(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")),
                    ("images", ("__slide_flow_platform__-template.bin", b"", "application/octet-stream")),
                ],
            )
        self.assertEqual(response.status_code, 200, response.text)
        schedule.assert_called_once()
        task_id = int(response.json()["task_id"])
        row = self.db.execute("SELECT params FROM tasks WHERE id = ?", (task_id,)).fetchone()
        params = json.loads(row["params"])
        self.assertEqual(params["import_target"], "templates")
        self.assertEqual(params["series"], "品牌系列")
        session = _load_resource_import_session_file(response.json()["session_id"])
        self.assertIsNotNone(session)
        self.assertEqual(session["import_target"], "templates")


if __name__ == "__main__":
    unittest.main()
