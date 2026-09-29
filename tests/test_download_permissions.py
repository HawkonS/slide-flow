from __future__ import annotations

import json
import re
import shutil
import sqlite3
import tempfile
import unittest
import zipfile
from io import BytesIO
from pathlib import Path
from unittest.mock import AsyncMock, patch

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation

from app.config import settings
from app.core import download_tasks
from app.core.permissions import SESSION_COOKIE, _auth_db_dep
from app.core.security import create_session_token
from app.routers import downloads
from app.routers.dependencies import db_dep, db_read_dep
from app.routers.shows import downloads as show_downloads
from app.services.downloads import cache
from app.services import files


class DownloadPermissionTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='slideflow-download-access-')
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name).resolve()
        self.assets = self.root / 'assets'
        self.output = self.root / 'downloads'
        self.assets.mkdir()
        self.output.mkdir()
        self.db_path = self.root / 'test.sqlite'
        for target, name, value in (
            (settings, 'root_dir', self.root),
            (settings, 'assets_dir', self.assets),
            (settings, 'downloads_dir', self.output),
            (settings, 'secret_key', 'download-test-only-signing-key'),
            (cache, '_DOWNLOAD_CACHE_DIR', self.output / 'cache'),
            (download_tasks, '_DOWNLOAD_TASKS_DIR', self.output / 'tasks'),
            (download_tasks, 'get_db', self.connect),
            (files, '_ALLOWED_FILE_DIRS', [self.assets, self.output]),
        ):
            patcher = patch.object(target, name, value)
            patcher.start()
            self.addCleanup(patcher.stop)
        self.db = self.connect()
        self.addCleanup(self.db.close)
        self.db.executescript('''
            CREATE TABLE users (id INTEGER PRIMARY KEY, role TEXT,
                session_version INTEGER DEFAULT 1, must_change_pwd INTEGER DEFAULT 0);
            INSERT INTO users(id, role) VALUES (1, 'user'), (2, 'user'),
                (3, 'admin'), (4, 'system_admin');
            CREATE TABLE shows (id INTEGER PRIMARY KEY, name TEXT, owner_id INTEGER,
                visibility_scope TEXT);
            INSERT INTO shows VALUES (10, 'Export', 1, 'public');
            CREATE TABLE resources (id INTEGER PRIMARY KEY, name TEXT, owner_id INTEGER,
                visibility_scope TEXT, current_version INTEGER);
            INSERT INTO resources VALUES (20, 'Public', 1, 'public', 2),
                (21, 'Private hidden', 1, 'private', 1);
            CREATE TABLE resource_versions (resource_id INTEGER, version_no INTEGER,
                ppt_path TEXT, png_path TEXT, font_names TEXT DEFAULT '[]',
                missing_fonts TEXT DEFAULT '[]');
            CREATE TABLE show_resources (show_id INTEGER, resource_id INTEGER,
                version_no INTEGER, sort_order INTEGER, is_hidden INTEGER);
            INSERT INTO show_resources VALUES (10, 20, 1, 0, 0), (10, 21, 1, 1, 1);
            CREATE TABLE tasks (id INTEGER PRIMARY KEY, owner_id INTEGER, task_type TEXT,
                status TEXT, params TEXT, result_data TEXT, progress INTEGER, total INTEGER);
            CREATE TABLE download_records (track_code TEXT, user_id INTEGER, show_id INTEGER,
                download_type TEXT, client_ip TEXT, downloaded_at TEXT);
            CREATE TABLE fonts (family_name TEXT, aliases TEXT, file_path TEXT);
        ''')
        for resource_id, version, color in ((20, 1, 'red'), (20, 2, 'blue'), (21, 1, 'green')):
            stem = f'{resource_id}-v{version}'
            png = self.assets / f'{stem}.png'
            ppt = self.assets / f'{stem}.pptx'
            Image.new('RGB', (64, 36), color).save(png)
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(0, 0, 1000000, 1000000).text = stem
            deck.save(ppt)
            self.db.execute('INSERT INTO resource_versions(resource_id, version_no, ppt_path, png_path) '
                            'VALUES (?, ?, ?, ?)', (resource_id, version, str(ppt), str(png)))
        self.db.commit()
        app = FastAPI()
        app.include_router(downloads.router)
        app.include_router(show_downloads.router)

        def db_override():
            yield self.db

        for dependency in (db_dep, db_read_dep, _auth_db_dep):
            app.dependency_overrides[dependency] = db_override
        self.client = TestClient(app)
        self.addCleanup(self.client.close)
        self.authenticate(2)

    def connect(self):
        db = sqlite3.connect(self.db_path, check_same_thread=False)
        db.row_factory = sqlite3.Row
        return db

    def authenticate(self, user_id, session_version=1):
        self.client.cookies.set(SESSION_COOKIE, create_session_token(
            user_id, settings.secret_key, session_version=session_version))

    def add_task(self, owner_id=2, download_type='zip', **overrides):
        params = {'show_id': 10, 'download_type': download_type, 'with_fonts': False,
                  'embed_fonts': False, 'track_code': 'test', 'session_version': 1}
        params.update(overrides)
        cursor = self.db.execute(
            "INSERT INTO tasks(owner_id, task_type, status, params) VALUES (?, 'download', 'pending', ?)",
            (owner_id, json.dumps(params)))
        self.db.commit()
        return cursor.lastrowid, params

    def generate(self, owner_id=2, download_type='zip', **overrides):
        task_id, params = self.add_task(owner_id, download_type, **overrides)
        result = download_tasks._generate_download_file_sync(task_id, params)
        self.db.execute("UPDATE tasks SET status = 'completed', result_data = ? WHERE id = ?",
                        (json.dumps(result), task_id))
        self.db.commit()
        return task_id, result

    def slide_count(self, response, kind):
        self.assertEqual(response.status_code, 200, response.text[:200] if response.status_code != 200 else '')
        if kind == 'pdf':
            return len(re.findall(rb'/Type\s*/Page\b', response.content))
        return len(Presentation(BytesIO(response.content)).slides)

    def test_shared_sync_cache_never_crosses_resource_permissions(self):
        for kind in ('pdf', 'pptx-images', 'pptx'):
            for first, second in ((1, 2), (2, 1)):
                with self.subTest(kind=kind, first=first):
                    shutil.rmtree(self.output / 'cache', ignore_errors=True)
                    for user_id in (first, second):
                        self.authenticate(user_id)
                        response = self.client.get(f'/api/shows/10/download/{kind}')
                        self.assertEqual(self.slide_count(response, kind), 2 if user_id == 1 else 1)

    def test_worker_exports_only_currently_visible_fixed_versions(self):
        _, result = self.generate()
        with zipfile.ZipFile(result['file_path']) as archive:
            ppt_names = [name for name in archive.namelist() if name.endswith('.pptx')]
            self.assertEqual(ppt_names, ['Public_v1.pptx'])
            self.assertEqual(archive.read(ppt_names[0]), (self.assets / '20-v1.pptx').read_bytes())

    def test_async_cache_is_isolated_and_records_permissions_on_cache_hits(self):
        self.generate(owner_id=1)
        _, result = self.generate()
        _, cached_result = self.generate()
        for output in (result, cached_result):
            with zipfile.ZipFile(output['file_path']) as archive:
                self.assertEqual([n for n in archive.namelist() if n.endswith('.pptx')], ['Public_v1.pptx'])
            self.assertEqual(output['resource_refs'], [[20, 1]])

    def test_worker_rechecks_show_and_session_before_using_cache(self):
        self.generate()
        task_id, params = self.add_task()
        for mutation, reset, expected_status in (
            ("UPDATE shows SET visibility_scope = 'private'", "UPDATE shows SET visibility_scope = 'public'", 403),
            ('UPDATE users SET session_version = 2 WHERE id = 2', 'UPDATE users SET session_version = 1 WHERE id = 2', 401),
            ('UPDATE users SET must_change_pwd = 1 WHERE id = 2', 'UPDATE users SET must_change_pwd = 0 WHERE id = 2', 403),
        ):
            with self.subTest(mutation=mutation):
                self.db.execute(mutation)
                self.db.commit()
                with patch.object(download_tasks, '_get_cached_download', return_value=None) as cached:
                    with self.assertRaises(HTTPException) as denied:
                        download_tasks._generate_download_file_sync(task_id, params)
                    self.assertEqual(denied.exception.status_code, expected_status)
                    cached.assert_not_called()
                self.db.execute(reset)
                self.db.commit()

    def test_completed_file_rechecks_show_and_included_resource_permissions(self):
        task_id, _ = self.generate()
        url = f'/api/downloads/{task_id}/file'
        self.assertEqual(self.client.get(url).status_code, 200)
        for table in ('shows', 'resources'):
            self.db.execute(f"UPDATE {table} SET visibility_scope = 'private'")
            self.db.commit()
            self.assertEqual(self.client.get(url).status_code, 403)
            self.db.execute(f"UPDATE {table} SET visibility_scope = 'public'")
            self.db.commit()

    def test_completed_file_rejects_revoked_session(self):
        task_id, _ = self.generate()
        self.db.execute('UPDATE users SET session_version = 2 WHERE id = 2')
        self.db.commit()
        self.assertEqual(self.client.get(f'/api/downloads/{task_id}/file').status_code, 401)
        self.authenticate(2, session_version=2)
        self.assertEqual(self.client.get(f'/api/downloads/{task_id}/file').status_code, 401)

    def test_removed_resource_is_still_checked_against_artifact_snapshot(self):
        self.db.execute("UPDATE resources SET visibility_scope = 'public' WHERE id = 21")
        self.db.commit()
        task_id, result = self.generate()
        self.assertEqual(result['resource_refs'], [[20, 1], [21, 1]])
        self.db.execute('DELETE FROM show_resources WHERE resource_id = 21')
        self.db.execute("UPDATE resources SET visibility_scope = 'private' WHERE id = 21")
        self.db.commit()
        self.assertEqual(self.client.get(f'/api/downloads/{task_id}/file').status_code, 403)

    def test_sync_cache_rechecks_permissions_for_the_same_user(self):
        for kind in ('pdf', 'pptx-images', 'pptx'):
            with self.subTest(kind=kind):
                self.db.execute("UPDATE resources SET visibility_scope = 'public'")
                self.db.commit()
                url = f'/api/shows/10/download/{kind}'
                self.assertEqual(self.slide_count(self.client.get(url), kind), 2)
                self.db.execute("UPDATE resources SET visibility_scope = 'private' WHERE id = 21")
                self.db.commit()
                self.assertEqual(self.slide_count(self.client.get(url), kind), 1)
                self.db.execute("UPDATE resources SET visibility_scope = 'private'")
                self.db.commit()
                self.assertEqual(self.client.get(url).status_code, 404)

    def test_five_formats_and_watermarks_preserve_authorized_content(self):
        for kind in ('pdf', 'pptx_images', 'pptx', 'pptx_pages', 'zip'):
            for watermark in ('', 'Test watermark'):
                with self.subTest(kind=kind, watermark=bool(watermark)):
                    with patch.object(download_tasks, '_get_cached_download', wraps=cache._get_cached_download) as cached:
                        _, result = self.generate(download_type=kind, user_watermark=watermark)
                        if watermark:
                            cached.assert_not_called()
                    self.assertEqual(result['resource_refs'], [[20, 1]])
                    self.assertEqual(result['watermark_applied'], bool(watermark))
                    data = Path(result['file_path']).read_bytes()
                    if kind == 'pdf':
                        self.assertEqual(len(re.findall(rb'/Type\s*/Page\b', data)), 1)
                    elif kind in ('pptx', 'pptx_images'):
                        self.assertEqual(len(Presentation(BytesIO(data)).slides), 1)
                    else:
                        with zipfile.ZipFile(BytesIO(data)) as archive:
                            self.assertEqual(len([n for n in archive.namelist() if n.endswith('.pptx')]), 1)

    def test_font_bundle_and_embedding_options_survive_authorization(self):
        for kind in ('pptx', 'zip'):
            with self.subTest(kind=kind):
                _, result = self.generate(download_type=kind, with_fonts=True)
                self.assertTrue(result['file_name'].endswith('_with_fonts.zip'))
                with zipfile.ZipFile(result['file_path']) as archive:
                    self.assertEqual(len([n for n in archive.namelist() if n.endswith('.pptx')]), 1)
        with patch.object(download_tasks, 'embed_fonts_in_pptx', side_effect=lambda path, db, fonts: path.read_bytes()) as embed:
            _, result = self.generate(download_type='pptx', embed_fonts=True)
            embed.assert_called_once()
            self.assertEqual(embed.call_args.args[2], [])
            self.assertEqual(result['file_name'], 'Export_embedded_fonts.pptx')
            self.assertEqual(len(Presentation(result['file_path']).slides), 1)

    def test_creation_binds_server_session_and_preserves_download_options(self):
        self.db.execute('UPDATE users SET session_version = 7 WHERE id = 2')
        self.db.commit()
        self.authenticate(2, session_version=7)
        with patch.object(download_tasks, 'execute_download_task', new_callable=AsyncMock):
            response = self.client.post('/api/downloads/create', json={
                'show_id': 10, 'download_type': 'pptx', 'embed_fonts': True,
                'watermark': 'Requested watermark', 'session_version': 999,
            })
        self.assertEqual(response.status_code, 200, response.text)
        task = self.db.execute('SELECT * FROM tasks WHERE id = ?', (response.json()['task_id'],)).fetchone()
        params = json.loads(task['params'])
        self.assertEqual(task['owner_id'], 2)
        self.assertEqual(params['session_version'], 7)
        self.assertTrue(params['embed_fonts'])
        self.assertEqual(params['user_watermark'], 'Requested watermark')

    def oss_materializer(self):
        sources = {}
        created = []
        for column, suffix in (('png_path', '.png'), ('ppt_path', '.pptx')):
            ref = f'oss://test/resources/public{suffix}'
            sources[ref] = self.assets / f'20-v1{suffix}'
            self.db.execute(f'UPDATE resource_versions SET {column} = ? WHERE resource_id = 20 AND version_no = 1', (ref,))
        self.db.commit()

        def materialize(ref):
            source = sources[ref]
            output = self.root / f'materialized-{len(created)}{source.suffix}'
            shutil.copy2(source, output)
            created.append(output)
            return output

        return materialize, created

    def test_sync_oss_exports_materialize_only_visible_content_and_clean_up(self):
        materialize, created = self.oss_materializer()
        with patch.object(files.oss_storage, 'materialize', side_effect=materialize):
            for kind in ('pdf', 'pptx-images', 'pptx', 'zip'):
                with self.subTest(kind=kind):
                    response = self.client.get(f'/api/shows/10/download/{kind}')
                    self.assertEqual(response.status_code, 200, response.text[:200] if response.status_code != 200 else '')
                    self.assertIn('no-store', response.headers['cache-control'])
                    if kind == 'zip':
                        with zipfile.ZipFile(BytesIO(response.content)) as archive:
                            self.assertEqual([n for n in archive.namelist() if n.endswith('.pptx')], ['Public_v1.pptx'])
                    else:
                        self.assertEqual(self.slide_count(response, kind), 1)
                    self.assertTrue(created)
                    self.assertFalse(any(path.exists() for path in created))

    def test_watermarked_oss_pdf_materializes_each_asset_once_and_cleans_up(self):
        materialize, created = self.oss_materializer()
        with patch.object(files.oss_storage, 'materialize', side_effect=materialize) as fetched:
            _, result = self.generate(download_type='pdf', user_watermark='Test watermark')
        self.assertEqual(result['resource_refs'], [[20, 1]])
        self.assertEqual(fetched.call_count, 1)
        self.assertFalse(any(path.exists() for path in created))

    def test_archive_members_cannot_contain_paths_from_user_names(self):
        self.db.execute('UPDATE resources SET name = ? WHERE id = 20', ('../../Folder\\payload',))
        self.db.execute('UPDATE shows SET name = ?', ('../../Folder\\Show',))
        self.db.commit()
        for path in ('zip', 'pptx?with_fonts=true'):
            with self.subTest(sync=path):
                response = self.client.get(f'/api/shows/10/download/{path}')
                self.assertEqual(response.status_code, 200)
                with zipfile.ZipFile(BytesIO(response.content)) as archive:
                    for name in archive.namelist():
                        if name.endswith('.pptx'):
                            self.assertNotRegex(name, r'[/\\]')
        for kind, options in (('zip', {}), ('pptx_pages', {}), ('pptx', {'with_fonts': True})):
            with self.subTest(async_type=kind):
                _, result = self.generate(download_type=kind, **options)
                with zipfile.ZipFile(result['file_path']) as archive:
                    for name in archive.namelist():
                        if name.endswith('.pptx'):
                            self.assertNotRegex(name, r'[/\\]')

    def test_operations_admin_cannot_download_another_users_artifact(self):
        task_id, _ = self.generate(owner_id=1)
        self.authenticate(3)
        self.assertEqual(self.client.get(f'/api/downloads/{task_id}/file').status_code, 403)
        self.authenticate(4)
        self.assertEqual(self.client.get(f'/api/downloads/{task_id}/file').status_code, 200)

    def test_artifacts_without_authorization_snapshot_must_be_regenerated(self):
        task_id, result = self.generate()
        result.pop('resource_refs', None)
        self.db.execute('UPDATE tasks SET result_data = ? WHERE id = ?', (json.dumps(result), task_id))
        self.db.commit()
        self.assertEqual(self.client.get(f'/api/downloads/{task_id}/file').status_code, 410)

    def test_task_creation_rejects_empty_authorized_export(self):
        self.db.execute("UPDATE resources SET visibility_scope = 'private'")
        self.db.commit()
        with patch.object(download_tasks, 'execute_download_task', new_callable=AsyncMock) as generate:
            response = self.client.post('/api/downloads/create', json={'show_id': 10, 'download_type': 'zip'})
            self.assertEqual(response.status_code, 404, response.text)
            generate.assert_not_called()
        self.assertEqual(self.db.execute('SELECT count(*) FROM tasks').fetchone()[0], 0)

    def test_completed_file_keeps_download_directory_boundary(self):
        task_id, result = self.generate()
        result['file_path'] = str(self.assets / '20-v1.pptx')
        self.db.execute('UPDATE tasks SET result_data = ? WHERE id = ?', (json.dumps(result), task_id))
        self.db.commit()
        self.assertEqual(self.client.get(f'/api/downloads/{task_id}/file').status_code, 410)


if __name__ == '__main__':
    unittest.main()
