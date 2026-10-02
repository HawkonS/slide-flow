"""Deletion contracts exercised against a migrated, isolated SQLite database."""
import sqlite3
import time
import unittest
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

from fastapi import HTTPException
from app.config import settings
from app.db import now_iso
from app.routers import presentation, show_shares, user_center
from app.routers.dependencies import db_dep
from app.routers.resources import files, mutations, queries, remarks
from app.routers.shows import catalog, versions
from app.services.resource_deletion import drain_file_gc, maintain_archives
from app.services.files import _init_allowed_file_dirs
from app.services.files import _init_allowed_file_dirs
from app.services.resources import _allocate_version_number
from app.services.shows import _collect_show_accessible_resources
from tests import test_resource_share_management as fixtures


class ResourceDeletionTests(unittest.TestCase):
    insert_user = fixtures.ResourceShareManagementTests.insert_user
    insert_resource = fixtures.ResourceShareManagementTests.insert_resource
    tearDown = fixtures.ResourceShareManagementTests.tearDown

    def setUp(self):
        fixtures.ResourceShareManagementTests.setUp(self)
        settings.resources_dir.mkdir(parents=True, exist_ok=True)
        _init_allowed_file_dirs()
        _init_allowed_file_dirs()
        for router in [queries.router, mutations.router, files.router, remarks.router,
                       catalog.router, versions.router, show_shares.router, presentation.router, user_center.router]:
            self.client.app.include_router(router)
        def transactional_db():
            try:
                yield self.db
            finally:
                if self.db.in_transaction:
                    self.db.rollback()
        self.client.app.dependency_overrides[db_dep] = transactional_db
        self.r = self.resource('待删除素材')
        self.other = self.resource('保留素材')

    def resource(self, name, owner=None):
        rid, _ = self.insert_resource(owner=owner or self.alice, name=name)
        self.db.execute('UPDATE resource_versions SET ppt_path = ?, png_path = ? WHERE resource_id = ?',
                        (str(settings.resources_dir / f'{rid}.pptx'), str(settings.resources_dir / f'{rid}.png'), rid))
        self.db.commit()
        return rid

    def version(self, rid, number):
        vid = self.db.execute('''INSERT INTO resource_versions
            (resource_id, version_no, ppt_path, png_path, common_remark_html, created_by, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?)''',
            (rid, number, str(settings.resources_dir / f'{rid}-v{number}.pptx'),
             str(settings.resources_dir / f'{rid}-v{number}.png'), f'<p>v{number}备注</p>', self.alice['id'], now_iso())).lastrowid
        self.db.execute('UPDATE resources SET current_version = ?, next_version_no = MAX(next_version_no, ?) WHERE id = ?', (number, number + 1, rid))
        self.db.commit()
        return vid

    def show(self, pages, *, owner=None, version=1, series=None, visibility='public'):
        owner = owner or self.alice
        sid = self.db.execute('''INSERT INTO shows
            (name, owner_id, subject, tags, status, visibility_scope, management_scope,
             series_id, version_no, change_note, updated_by, created_at, updated_at)
            VALUES (?, ?, '', '', 'active', ?, 'private', ?, ?, '', ?, ?, ?)''',
            ('机密放映' if owner['id'] == self.bob['id'] else '引用放映', owner['id'], visibility,
             series or str(uuid4()), version, owner['id'], now_iso(), now_iso())).lastrowid
        for order, value in enumerate(pages):
            rid, number, hidden = (value, 1, 0) if isinstance(value, int) else value
            self.db.execute('INSERT INTO show_resources VALUES (?, ?, ?, ?, ?)', (sid, rid, number, order, hidden))
        self.db.commit()
        return sid

    def preview(self, ids=None, scope='all'):
        result = self.client.post('/api/resources/delete-preview', json={'resource_ids': ids or [self.r], 'scope': scope})
        self.assertEqual(result.status_code, 200, result.text)
        return result.json()

    def delete(self, *, action='remove', scope='all', ids=None, preview=None, rollback=False):
        ids = ids or [self.r]
        preview = preview or self.preview(ids, scope)
        body = {'reference_action': action, 'confirmation_token': preview['confirmation_token']}
        if rollback:
            return self.client.post(f'/api/resources/{ids[0]}/versions/rollback', json=body)
        if len(ids) > 1:
            return self.client.request('DELETE', '/api/resources/batch', json={**body, 'resource_ids': ids, 'scope': scope})
        return self.client.request('DELETE', f'/api/resources/{ids[0]}?scope={scope}', json=body)

    def pages(self, sid):
        return [tuple(r) for r in self.db.execute('SELECT resource_id, version_no, is_hidden FROM show_resources WHERE show_id = ? ORDER BY sort_order', (sid,))]

    def share(self, sid):
        response = self.client.post(f'/api/shows/{sid}/share-links', json={'expires_in_days': 7})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def test_unreferenced_delete_keeps_unrelated_legacy_empty_show(self):
        empty = self.show([])
        self.assertFalse(self.preview()['has_references'])
        result = self.client.delete(f'/api/resources/{self.r}')
        self.assertEqual(result.status_code, 200, result.text)
        self.assertIsNone(self.db.execute('SELECT * FROM resources WHERE id = ?', (self.r,)).fetchone())
        self.assertEqual(self.client.get(f'/api/shows/{empty}').status_code, 200)

    def test_preview_deduplicates_batch_and_includes_hidden_historical_versions(self):
        self.version(self.r, 2)
        old = self.show([(self.r, 1, 1), self.other], version=1, series='history')
        new = self.show([(self.r, 2, 0)], version=2, series='history')
        result = self.preview([self.r, self.other, self.r])
        self.assertEqual((result['show_count'], result['removed_pages'], result['empty_show_count']), (2, 3, 2))
        self.assertEqual({s['id'] for s in result['shows']}, {old, new})
        self.assertTrue(next(s for s in result['shows'] if s['id'] == old)['references'][0]['hidden'])
        latest = self.preview(scope='latest')
        self.assertEqual([s['id'] for s in latest['shows']], [new])

    def test_preview_includes_owner_for_visible_show(self):
        sid = self.show([self.r], owner=self.bob, visibility='public')
        preview = self.preview()
        show = next(item for item in preview['shows'] if item['id'] == sid)
        self.assertEqual(show['owner']['name'], self.bob['name'])
        self.assertEqual(show['owner']['username'], self.bob['username'])

    def test_full_scope_targets_only_active_versions(self):
        self.version(self.r, 2)
        archived = self.db.execute(
            'SELECT id FROM resource_versions WHERE resource_id = ? AND version_no = 1',
            (self.r,),
        ).fetchone()['id']
        self.db.execute(
            'UPDATE resource_versions SET deleted_at = ? WHERE id = ?',
            (now_iso(), archived),
        )
        self.db.commit()
        preview = self.preview(scope='all')
        target = preview['targets'][0]
        self.assertEqual(target['version_nos'], [2])
        self.assertEqual(target['selected_version_count'], 1)

    def test_references_require_strategy_and_valid_confirmation(self):
        sid = self.show([self.r])
        self.assertEqual(self.client.delete(f'/api/resources/{self.r}').status_code, 409)
        self.assertEqual(self.client.request('DELETE', f'/api/resources/{self.r}', json={'reference_action': 'remove'}).status_code, 409)
        preview = self.preview()
        bad = {**preview, 'confirmation_token': preview['confirmation_token'] + 'bad'}
        self.assertEqual(self.delete(preview=bad).status_code, 409)
        with patch('app.services.resource_deletion.time.time', return_value=time.time() + 901):
            self.assertEqual(self.delete(preview=preview).status_code, 409)
        self.assertEqual(self.pages(sid), [(self.r, 1, 0)])

    def test_preserve_keeps_files_notes_download_inputs_and_shares(self):
        from PIL import Image
        from pptx import Presentation
        row = self.db.execute('SELECT * FROM resource_versions WHERE resource_id = ?', (self.r,)).fetchone()
        deck = Presentation(); deck.slides.add_slide(deck.slide_layouts[6]); deck.save(row['ppt_path'])
        Image.new('RGB', (32, 18), 'blue').save(row['png_path'])
        sid = self.show([self.r, self.other])
        self.db.execute('INSERT INTO personal_remarks (resource_id, version_id, user_id, content_html, updated_at) VALUES (?, ?, ?, ?, ?)', (self.r, row['id'], self.alice['id'], '<p>个人原备注</p>', now_iso()))
        self.db.execute('INSERT INTO show_remarks (show_id, resource_id, user_id, content_html, updated_at) VALUES (?, ?, ?, ?, ?)', (sid, self.r, self.alice['id'], '<p>放映原备注</p>', now_iso()))
        self.db.commit()
        link = self.share(sid)
        resource_link = self.client.post(f'/api/resources/{self.r}/share-links', json={'expires_in_days': 7}).json()
        result = self.delete(action='preserve')
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()['removed_pages'], 0)
        self.assertEqual(self.client.get(f'/api/resources/{self.r}').status_code, 404)
        archived = self.client.get(f'/api/resources/{self.r}?include_archived=true').json()['resource']
        self.assertTrue(archived['versions'][0]['archived'])
        show = self.client.get(f'/api/shows/{sid}').json()['show']
        self.assertTrue(show['resources'][0]['accessible'])
        self.assertEqual(show['resources'][0]['version_no'], 1)
        self.assertEqual(self.client.get(f'/api/resources/{self.r}/preview?version_id={row["id"]}').status_code, 200)
        self.assertTrue(Path(row['ppt_path']).exists())
        self.assertEqual(_collect_show_accessible_resources(self.db, sid, self.alice)[0]['resource_id'], self.r)
        self.assertEqual(self.client.get(f'/api/resources/{self.r}/personal-remark?version_id={row["id"]}').json()['content_html'], '<p>个人原备注</p>')
        self.assertEqual(self.db.execute('SELECT content_html FROM show_remarks WHERE show_id = ?', (sid,)).fetchone()[0], '<p>放映原备注</p>')
        shared = self.client.get('/api/show-shares/' + link['token'])
        self.assertEqual(shared.status_code, 200, shared.text)
        self.assertIn(self.r, [p['id'] for p in shared.json()['show']['resources']])
        self.assertEqual(self.client.get('/api/resource-shares/' + resource_link['token']).status_code, 404)
        self.assertNotIn(self.r, self.client.get('/api/resources/ids').json()['ids'])
        self.assertNotIn(self.r, self.client.get('/api/resources/pick-ids').json()['ids'])
        self.assertEqual(self.client.post(f'/api/resources/{self.r}/common-remark', json={'content_html': '修改'}).status_code, 404)
        self.assertEqual(self.client.put(f'/api/resources/{self.r}/personal-remark', json={'content_html': '修改', 'version_id': row['id']}).status_code, 404)

    def test_latest_preserve_and_rollback_never_reuse_version_numbers(self):
        vid = self.version(self.r, 2)
        sid = self.show([(self.r, 2, 0)])
        result = self.delete(action='preserve', scope='latest', rollback=True)
        self.assertEqual(result.status_code, 200, result.text)
        resource = result.json()['resource']
        self.assertEqual((resource['current_version'], resource['version_count']), (1, 1))
        self.assertEqual(self.pages(sid), [(self.r, 2, 0)])
        self.assertEqual(self.client.post(f'/api/resources/{self.r}/common-remark', json={'content_html': '不应写入', 'apply_scope': 'selected', 'version_id': vid}).status_code, 404)
        number = _allocate_version_number(self.db, self.r)
        self.assertEqual(number, 3)
        self.db.commit()
        self.version(self.r, number)
        self.assertEqual(self.pages(sid), [(self.r, 2, 0)])
        self.assertEqual(self.client.post(f'/api/resources/{self.r}/versions/rollback').status_code, 200)
        self.assertEqual(_allocate_version_number(self.db, self.r), 4)
        self.db.rollback()

    def test_remove_latest_deletes_only_newly_empty_show_versions(self):
        self.version(self.r, 2)
        old = self.show([self.r], series='same', version=1)
        new = self.show([(self.r, 2, 1)], series='same', version=2)
        mixed = self.show([(self.r, 2, 0), self.other])
        result = self.delete(scope='latest')
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()['deleted_show_ids'], [new])
        self.assertEqual(self.pages(old), [(self.r, 1, 0)])
        self.assertEqual(self.pages(mixed), [(self.other, 1, 0)])
        self.assertEqual(self.client.get(f'/api/shows/{new}').status_code, 404)
        self.assertEqual(self.db.execute('SELECT sort_order FROM show_resources WHERE show_id = ?', (mixed,)).fetchone()[0], 0)

    def test_full_removal_also_cleans_dangling_version_references(self):
        sid = self.show([(self.r, 999, 1)])
        self.assertEqual(self.preview()['removed_pages'], 1)
        result = self.delete()
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()['deleted_show_ids'], [sid])

    def test_invisible_unmanageable_show_exposes_only_count_and_allows_preserve(self):
        sid = self.show([self.r], owner=self.bob, visibility='private')
        preview = self.preview()
        self.assertEqual(preview['shows'], [])
        self.assertEqual(preview['restricted_show_count'], 1)
        self.assertEqual(preview['available_actions'], ['preserve'])
        self.assertNotIn('机密放映', str(preview))
        self.assertEqual(self.delete(preview=preview).status_code, 403)
        self.assertEqual(self.delete(preview=preview, action='preserve').status_code, 200)
        self.assertEqual(self.pages(sid), [(self.r, 1, 0)])

    def test_stale_confirmation_rejects_new_reference_version_and_permission_change(self):
        for change in ['reference', 'version', 'permission']:
            with self.subTest(change=change):
                rid = self.resource(change)
                sid = self.show([rid])
                preview = self.preview([rid])
                if change == 'reference': self.show([rid])
                elif change == 'version': self.version(rid, 2)
                else:
                    self.db.execute('UPDATE shows SET owner_id = ? WHERE id = ?', (self.bob['id'], sid)); self.db.commit()
                self.assertEqual(self.delete(ids=[rid], preview=preview).status_code, 409)
                self.assertIsNone(self.db.execute('SELECT deleted_at FROM resources WHERE id = ?', (rid,)).fetchone()[0])
                self.assertEqual(self.pages(sid), [(rid, 1, 0)])

    def test_batch_authorization_and_database_failure_are_atomic(self):
        sid = self.show([self.r, self.other])
        forbidden = self.resource('无权删除', self.bob)
        response = self.client.request('DELETE', '/api/resources/batch', json={'resource_ids': [self.r, forbidden]})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(len(self.pages(sid)), 2)
        preview = self.preview([self.r, self.other])
        self.db.execute(f'''CREATE TRIGGER reject_second_delete BEFORE UPDATE OF deleted_at ON resources
            WHEN NEW.id = {self.other} BEGIN SELECT RAISE(ABORT, 'injected failure'); END''')
        self.db.commit()
        with self.assertRaises(sqlite3.IntegrityError):
            self.delete(ids=[self.r, self.other], preview=preview)
        self.assertEqual(len(self.pages(sid)), 2)
        self.assertEqual(self.db.execute('SELECT count(*) FROM resources WHERE deleted_at IS NOT NULL').fetchone()[0], 0)

    def test_batch_preserve_and_repeat_submission_do_not_remove_pages(self):
        sid = self.show([self.r, self.other])
        preview = self.preview([self.r, self.other])
        self.assertEqual(self.delete(ids=[self.r, self.other], action='preserve', preview=preview).status_code, 200)
        self.assertEqual(self.delete(ids=[self.r, self.other], action='preserve', preview=preview).status_code, 404)
        self.assertEqual(len(self.pages(sid)), 2)

    def test_share_only_references_are_cleaned_and_empty_snapshot_revoked(self):
        sid = self.show([self.r])
        single = self.share(sid)
        self.db.execute('INSERT INTO show_resources VALUES (?, ?, 1, 1, 0)', (sid, self.other)); self.db.commit()
        mixed = self.share(sid)
        self.db.execute('DELETE FROM show_resources WHERE show_id = ? AND resource_id = ?', (sid, self.r)); self.db.commit()
        preview = self.preview()
        self.assertEqual((preview['removed_pages'], preview['share_count'], preview['empty_show_count']), (0, 2, 0))
        result = self.delete(preview=preview)
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()['deleted_show_ids'], [])
        self.assertEqual(self.client.get('/api/show-shares/' + single['token']).status_code, 404)
        response = self.client.get('/api/show-shares/' + mixed['token'])
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual([p['id'] for p in response.json()['show']['resources']], [self.other])

    def test_archive_is_retained_for_share_snapshot_until_last_reference_expires(self):
        sid = self.show([self.r])
        link = self.share(sid)
        self.assertEqual(self.delete(action='preserve').status_code, 200)
        self.db.execute('DELETE FROM show_resources WHERE show_id = ?', (sid,)); self.db.commit()
        maintain_archives(self.db)
        self.assertEqual(self.client.get('/api/show-shares/' + link['token']).status_code, 200)
        self.db.execute("UPDATE show_share_tokens SET expires_at = '2000-01-01T00:00:00Z'"); self.db.commit()
        maintain_archives(self.db)
        self.assertIsNone(self.db.execute('SELECT * FROM resources WHERE id = ?', (self.r,)).fetchone())

    def test_file_cleanup_failure_is_retryable_without_failing_committed_delete(self):
        path = settings.resources_dir / f'{self.r}.pptx'; path.write_bytes(b'original')
        with patch('app.services.resource_deletion._delete_resource_files', side_effect=OSError('storage offline')):
            with self.assertLogs('app.services.resource_deletion', level='ERROR'):
                response = self.delete()
        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(path.exists())
        self.assertEqual(self.db.execute('SELECT count(*) FROM resource_file_gc').fetchone()[0], 1)
        drain_file_gc(self.db)
        self.assertFalse(path.exists())
        self.assertEqual(self.db.execute('SELECT count(*) FROM resource_file_gc').fetchone()[0], 0)

    def test_missing_version_is_explicit_and_cleanup_requires_fresh_confirmation(self):
        sid = self.show([(self.r, 77, 0), self.other])
        resource = self.client.get(f'/api/shows/{sid}').json()['show']['resources'][0]
        self.assertEqual(resource['unavailable_reason'], 'missing_version')
        with self.assertRaises(HTTPException) as error:
            _collect_show_accessible_resources(self.db, sid, self.alice)
        self.assertEqual(error.exception.status_code, 409)
        url = f'/api/shows/{sid}/cleanup-missing-resources'
        preview = self.client.post(url, json={'preview': True}).json()
        self.assertEqual((preview['missing_count'], preview['remaining_pages'], preview['will_delete']), (1, 1, False))
        self.assertEqual(self.client.post(url, json={'preview': False}).status_code, 409)
        response = self.client.post(url, json={'preview': False, 'confirmation_token': preview['confirmation_token']})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(self.pages(sid), [(self.other, 1, 0)])

    def test_missing_resource_and_empty_show_can_be_explicitly_cleaned(self):
        sid = self.show([])
        self.db.execute('PRAGMA foreign_keys = OFF')
        self.db.execute('INSERT INTO show_resources VALUES (?, 999999, 1, 0, 0)', (sid,)); self.db.commit()
        self.db.execute('PRAGMA foreign_keys = ON')
        page = self.client.get(f'/api/shows/{sid}').json()['show']['resources'][0]
        self.assertEqual(page['unavailable_reason'], 'missing_resource')
        url = f'/api/shows/{sid}/cleanup-missing-resources'
        preview = self.client.post(url, json={}).json()
        self.assertTrue(preview['will_delete'])
        result = self.client.post(url, json={'preview': False, 'confirmation_token': preview['confirmation_token']})
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()['deleted_show_ids'], [sid])

    def test_preserved_pages_can_be_copied_or_iterated_but_not_added_anew(self):
        sid = self.show([self.r])
        self.assertEqual(self.delete(action='preserve').status_code, 200)
        copied = self.client.post(f'/api/shows/{sid}/duplicate', json={'name': '副本'})
        self.assertEqual(copied.status_code, 200, copied.text)
        self.assertEqual(self.pages(copied.json()['show']['id']), [(self.r, 1, 0)])
        iteration = self.client.post(f'/api/shows/{sid}/iterate', json={'resource_ids': [self.r, self.other]})
        self.assertEqual(iteration.status_code, 200, iteration.text)
        self.assertEqual(self.pages(iteration.json()['show']['id']), [(self.r, 1, 0), (self.other, 1, 0)])
        fresh = self.show([self.other])
        self.assertEqual(self.client.post(f'/api/shows/{fresh}/resources/append', json={'resource_id': self.r}).status_code, 404)
        self.assertEqual(self.client.post('/api/shows', json={'name': '新建', 'status': 'active', 'resource_ids': [self.r]}).status_code, 404)

    def test_cleanup_permissions_and_last_version_rollback(self):
        sid = self.show([(self.r, 99, 0)], owner=self.bob)
        self.assertEqual(self.client.post(f'/api/shows/{sid}/cleanup-missing-resources', json={}).status_code, 403)
        self.assertEqual(self.client.post(f'/api/resources/{self.r}/versions/rollback').status_code, 400)
        self.assertEqual(self.pages(sid), [(self.r, 99, 0)])
