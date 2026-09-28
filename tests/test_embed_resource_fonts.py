from __future__ import annotations

from io import BytesIO
from pathlib import Path
import tempfile
import sqlite3
import unittest
import zipfile
import xml.etree.ElementTree as ET
from unittest.mock import patch

from fontTools.fontBuilder import FontBuilder
from fontTools.pens.t2CharStringPen import T2CharStringPen
from fontTools.pens.ttGlyphPen import TTGlyphPen
from fontTools.ttLib import TTFont
from pptx import Presentation

from app.services.downloads.embed_fonts import (
    FONT_REL_TYPE,
    FontEmbeddingError,
    FontRow,
    embed_fonts_in_pptx,
    prepare_embedded_font_cache,
)
from app.config import settings


P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _make_font(path: Path, *, fs_type: int = 0) -> None:
    builder = FontBuilder(1000, isTTF=True)
    glyph_order = [".notdef", "space"]
    builder.setupGlyphOrder(glyph_order)
    builder.setupCharacterMap({32: "space"})
    pen = TTGlyphPen(None)
    empty_glyph = pen.glyph()
    builder.setupGlyf({".notdef": empty_glyph, "space": empty_glyph})
    builder.setupHorizontalMetrics({".notdef": (500, 0), "space": (250, 0)})
    builder.setupHorizontalHeader(ascent=800, descent=-200)
    builder.setupNameTable(
        {
            "familyName": "Test Sans",
            "styleName": "Regular",
            "uniqueFontIdentifier": "Test Sans Regular",
            "fullName": "Test Sans Regular",
            "psName": "TestSans-Regular",
        }
    )
    builder.setupOS2(
        sTypoAscender=800,
        sTypoDescender=-200,
        usWinAscent=800,
        usWinDescent=200,
        fsType=fs_type,
    )
    builder.setupPost()
    builder.setupMaxp()
    builder.save(path)


def _make_cff_font(path: Path) -> None:
    builder = FontBuilder(1000, isTTF=False)
    glyph_order = [".notdef", "space", "A"]
    builder.setupGlyphOrder(glyph_order)
    builder.setupCharacterMap({32: "space", 65: "A"})
    char_strings = {}
    for glyph_name, width in ((".notdef", 500), ("space", 250), ("A", 600)):
        pen = T2CharStringPen(width, None)
        if glyph_name == "A":
            pen.moveTo((0, 0))
            pen.lineTo((300, 700))
            pen.lineTo((600, 0))
            pen.closePath()
        char_strings[glyph_name] = pen.getCharString()
    builder.setupHorizontalMetrics({".notdef": (500, 0), "space": (250, 0), "A": (600, 0)})
    builder.setupHorizontalHeader(ascent=800, descent=-200)
    builder.setupNameTable(
        {
            "familyName": "Test CFF",
            "styleName": "Regular",
            "uniqueFontIdentifier": "Test CFF Regular",
            "fullName": "Test CFF Regular",
            "psName": "TestCFF-Regular",
        }
    )
    builder.setupOS2(
        sTypoAscender=800,
        sTypoDescender=-200,
        usWinAscent=800,
        usWinDescent=200,
    )
    builder.setupPost()
    builder.setupMaxp()
    builder.setupCFF(
        "TestCFF-Regular",
        {"FullName": "Test CFF Regular", "FamilyName": "Test CFF", "Weight": "Regular"},
        char_strings,
        {"nominalWidthX": 0, "defaultWidthX": 0},
    )
    builder.save(path)


def _make_pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text = slide.shapes.add_textbox(0, 0, 1000, 400)
    text.text_frame.text = "Embedded font test"
    presentation.save(path)


class EmbeddedFontDownloadTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="slideflow-embedded-font-")
        self.root = Path(self.temporary.name)
        self.font_path = self.root / "test.ttf"
        self.ppt_path = self.root / "source.pptx"
        _make_font(self.font_path)
        _make_pptx(self.ppt_path)
        self.row = FontRow("Test Sans", ("Test Sans",), self.font_path)

    def tearDown(self):
        self.temporary.cleanup()

    def test_embeds_font_data_and_relationship_in_presentation_package(self):
        with patch(
            "app.services.downloads.embed_fonts._font_rows",
            return_value=({"testsans": self.row}, [self.row]),
        ):
            content = embed_fonts_in_pptx(self.ppt_path, None, ["Test Sans"])

        with zipfile.ZipFile(BytesIO(content)) as package:
            font_data = package.read("ppt/fonts/font1.fntdata")
            self.assertEqual(int.from_bytes(font_data[:4], "little"), len(font_data))
            font_payload_size = int.from_bytes(font_data[4:8], "little")
            self.assertGreater(font_payload_size, 0)
            self.assertEqual(font_data[34:36], b"LP")
            self.assertEqual(font_data[-font_payload_size:-font_payload_size + 4], b"\x00\x01\x00\x00")
            presentation = ET.fromstring(package.read("ppt/presentation.xml"))
            self.assertEqual(presentation.attrib["embedTrueTypeFonts"], "1")
            embedded = presentation.find(f"{{{P_NS}}}embeddedFontLst/{{{P_NS}}}embeddedFont")
            self.assertIsNotNone(embedded)
            self.assertEqual(embedded.find(f"{{{P_NS}}}font").attrib["typeface"], "Test Sans")
            font_rel_id = embedded.find(f"{{{P_NS}}}regular").attrib[f"{{{R_NS}}}id"]
            relationships = ET.fromstring(package.read("ppt/_rels/presentation.xml.rels"))
            font_rel = next(rel for rel in relationships if rel.attrib.get("Id") == font_rel_id)
            self.assertEqual(font_rel.attrib["Type"], FONT_REL_TYPE)
            self.assertEqual(font_rel.attrib["Target"], "fonts/font1.fntdata")
            child_order = [child.tag.rsplit("}", 1)[-1] for child in presentation]
            self.assertLess(child_order.index("embeddedFontLst"), child_order.index("defaultTextStyle"))
            content_types = ET.fromstring(package.read("[Content_Types].xml"))
            self.assertTrue(any(node.attrib.get("Extension") == "fntdata" for node in content_types))

        self.assertEqual(len(Presentation(BytesIO(content)).slides), 1)

    def test_standard_library_fonts_can_embed_even_when_source_marks_restriction(self):
        _make_font(self.font_path, fs_type=0x0002)
        with patch(
            "app.services.downloads.embed_fonts._font_rows",
            return_value=({"testsans": self.row}, [self.row]),
        ):
            content = embed_fonts_in_pptx(self.ppt_path, None, ["Test Sans"])

        with zipfile.ZipFile(BytesIO(content)) as package:
            self.assertIn("ppt/fonts/font1.fntdata", package.namelist())
            font_data = package.read("ppt/fonts/font1.fntdata")
            # The generated EOT payload clears source-level restriction bits.
            self.assertEqual(int.from_bytes(font_data[32:34], "little") & 0x0202, 0)

    def test_converts_cff_font_to_truetype_payload(self):
        cff_path = self.root / "test.otf"
        _make_cff_font(cff_path)
        row = FontRow("Test CFF", ("Test CFF",), cff_path)
        with patch(
            "app.services.downloads.embed_fonts._font_rows",
            return_value=({"testcff": row}, [row]),
        ):
            content = embed_fonts_in_pptx(self.ppt_path, None, ["Test CFF"])

        with zipfile.ZipFile(BytesIO(content)) as package:
            font_data = package.read("ppt/fonts/font1.fntdata")
            font_payload_size = int.from_bytes(font_data[4:8], "little")
            payload = font_data[-font_payload_size:]
            self.assertEqual(payload[:4], b"\x00\x01\x00\x00")
            converted = TTFont(BytesIO(payload), lazy=False)
            try:
                self.assertNotIn("CFF ", converted)
                self.assertNotIn("CFF2", converted)
                self.assertIn("glyf", converted)
            finally:
                converted.close()

    def test_reuses_persisted_cache_without_converting_again(self):
        db = sqlite3.connect(":memory:")
        db.row_factory = sqlite3.Row
        db.executescript(
            """
            CREATE TABLE fonts (
                id INTEGER PRIMARY KEY,
                family_name TEXT NOT NULL,
                aliases TEXT NOT NULL,
                file_path TEXT NOT NULL
            );
            CREATE TABLE embedded_font_cache (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                font_id INTEGER NOT NULL,
                source_sha256 TEXT NOT NULL,
                face_key TEXT NOT NULL,
                aliases TEXT NOT NULL DEFAULT '[]',
                variant TEXT NOT NULL,
                cache_ref TEXT NOT NULL,
                status TEXT NOT NULL,
                converter_version TEXT NOT NULL,
                error_message TEXT,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                UNIQUE(font_id, source_sha256, face_key, converter_version)
            );
            """
        )
        db.execute(
            "INSERT INTO fonts (id, family_name, aliases, file_path) VALUES (1, ?, ?, ?)",
            ("Test Sans", '["Test Sans"]', str(self.font_path)),
        )
        db.commit()
        with patch.object(settings, "fonts_dir", self.root), patch.object(settings, "storage_backend", "local"):
            result = prepare_embedded_font_cache(db, 1)
            self.assertEqual(result["ready"], 1)
            # The cache is independent from the source file after preparation.
            self.font_path.unlink()
            with patch(
                "app.services.downloads.embed_fonts._font_data",
                side_effect=AssertionError("cache miss unexpectedly converted the font"),
            ):
                content = embed_fonts_in_pptx(self.ppt_path, db, ["Test Sans"])
        self.assertGreater(len(content), 0)
        db.close()

    def test_reports_missing_fonts(self):
        with patch(
            "app.services.downloads.embed_fonts._font_rows",
            return_value=({}, []),
        ):
            with self.assertRaisesRegex(FontEmbeddingError, "字体库中缺少"):
                embed_fonts_in_pptx(self.ppt_path, None, ["Not Installed"])


if __name__ == "__main__":
    unittest.main()
